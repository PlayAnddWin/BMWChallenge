"""Versioned, source-preserving conversion for BMW workbooks and PDFs.

The converter deliberately preserves uncertain spreadsheet cells as native
source records. It does not guess respondent counts, powertrain splits, source
channel semantics, or mappings between study questions and customer topics.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


CONVERSION_VERSION = "1.2.1"
SUPPORTED_EXCEL = {".xlsx", ".xlsm"}
SUPPORTED_PDF = {".pdf"}
DEFAULT_REQUIRED_KINDS = ("feedback", "survey", "sales", "document")
DEFAULT_REQUIRED_DOCUMENT_KINDS = ("option_catalog", "corporate_report")


@dataclass
class ConversionBundle:
    """In-memory output of a deterministic conversion run."""

    source_files: list[dict[str, Any]] = field(default_factory=list)
    source_records: list[dict[str, Any]] = field(default_factory=list)
    feedback_comments: list[dict[str, Any]] = field(default_factory=list)
    feedback_annotations: list[dict[str, Any]] = field(default_factory=list)
    facts: list[dict[str, Any]] = field(default_factory=list)
    document_chunks: list[dict[str, Any]] = field(default_factory=list)
    issues: list[dict[str, Any]] = field(default_factory=list)
    manifest: dict[str, Any] = field(default_factory=dict)


class ConversionError(RuntimeError):
    """Raised when a source cannot be safely converted."""


def json_value(value: Any) -> Any:
    """Convert common workbook values into stable JSON-compatible values."""
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def stable_id(prefix: str, *parts: Any) -> str:
    material = "\x1f".join("" if item is None else str(item) for item in parts)
    return f"{prefix}-{hashlib.sha256(material.encode('utf-8')).hexdigest()}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "feedback_id": ("id", "feedbackid", "customerfeedbackid", "commentid", "reviewid", "recordid"),
    "text": ("text", "feedback", "feedbacktext", "comment", "commenttext", "customerfeedback",
             "customerfeedbacktext", "reviewtext", "originaltext", "verbatim"),
    "feedback_type": ("feedbacktype", "type", "annotationtype", "sentimenttype"),
    "model": ("model", "vehiclemodel", "derivative", "derivateecode", "series", "vehicleline"),
    "country": ("country", "countrycode", "market", "marketcountry"),
    "region": ("region", "commercialregion", "salesregion"),
    "powertrain": ("powertrain", "enginetype", "engine", "drive", "drivetrain"),
    "source_label": ("source", "sourcelabel", "sourcechannel", "sourceletter"),
    "vfc2": ("vfc2", "vfc2name", "vfclevel2", "vfclevel2name", "vfclevel2category",
             "vfc2category", "topic", "topiccategory"),
    "vfc3": ("vfc3", "vfc3name", "vfclevel3", "vfclevel3name", "vfclevel3category",
             "vfc3category", "subtopic", "issuecategory"),
    "year": ("year", "forecastyear", "fiscalyear"),
    "volume": ("volume", "salesvolume", "units", "vehicles", "forecastvolume"),
    "respondent_n": ("respondentn", "respondents", "respondentcount", "samplesize", "nrespondents", "n"),
    "scale": ("scale", "ratingscale", "scorescale", "responsescale", "measurementscale"),
    "unit": ("unit", "measurementunit", "unitofmeasure"),
}


def _alias_value(headers: Sequence[str], values: Sequence[Any], field_name: str) -> Any:
    accepted = set(HEADER_ALIASES[field_name])
    for index, header in enumerate(headers):
        if normalize_header(header) in accepted and index < len(values):
            return values[index]
    return None


def _header_map(values: Sequence[Any]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        label = str(value).strip() if value is not None else ""
        label = label or f"column_{index}"
        seen[label] = seen.get(label, 0) + 1
        if seen[label] > 1:
            label = f"{label}__{seen[label]}"
        result.append(label)
    return result


def _row_payload(headers: Sequence[str], values: Sequence[Any]) -> dict[str, Any]:
    return {
        "columns": [
            {"header": headers[i] if i < len(headers) else f"column_{i + 1}", "value": json_value(value)}
            for i, value in enumerate(values)
        ],
        "values_by_header": {
            headers[i] if i < len(headers) else f"column_{i + 1}": json_value(value)
            for i, value in enumerate(values)
        },
    }


def _nonempty(values: Iterable[Any]) -> bool:
    return any(value is not None and str(value).strip() for value in values)


def _excel_module() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on runtime install
        raise ConversionError("Excel conversion requires openpyxl; install pipeline/requirements-data.txt") from exc
    return openpyxl


def _pdf_reader(path: Path) -> Any:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on runtime install
        raise ConversionError("PDF conversion requires pypdf; install pipeline/requirements-data.txt") from exc
    return PdfReader(path)


def classify_source(path: Path) -> tuple[str, str | None]:
    name = path.stem.casefold()
    compact = normalize_header(name)
    if path.suffix.casefold() in SUPPORTED_PDF:
        if any(token in compact for token in ("option", "equipment", "catalog", "catalogue", "price")):
            return "document", "option_catalog"
        if "report" in compact or "annual" in compact or "strategy" in compact:
            return "document", "corporate_report"
        return "document", "document"
    if path.suffix.casefold() in SUPPORTED_EXCEL:
        if any(token in compact for token in ("feedback", "explorer", "hackathon")):
            return "feedback", None
        if any(token in compact for token in ("study", "studies", "survey", "customer")):
            return "survey", None
        if any(token in compact for token in ("sales", "volume", "forecast")):
            return "sales", None
        return "workbook", None
    return "unsupported", None


def _source_file(path: Path, kind: str, document_kind: str | None) -> dict[str, Any]:
    digest = sha256_file(path)
    return {
        "id": f"SRC-{digest}",
        "filename": path.name,
        "sha256": digest,
        "source_version": digest,
        "kind": kind,
        "document_kind": document_kind,
        "conversion_version": CONVERSION_VERSION,
        "size_bytes": path.stat().st_size,
    }


def _row_source_record(source: Mapping[str, Any], sheet: str, row_num: int,
                       headers: Sequence[str], values: Sequence[Any]) -> dict[str, Any]:
    record_id = stable_id("SR", source["id"], sheet, row_num)
    return {
        "id": record_id,
        "source_file_id": source["id"],
        "source_version": source["source_version"],
        "locator": {"filename": source["filename"], "sheet": sheet, "row": row_num},
        "raw_payload": _row_payload(headers, values),
        "original_record_id": None,
    }


def _find_feedback_sheet(workbook: Any) -> tuple[Any, int, list[str]]:
    candidates = sorted(workbook.worksheets, key=lambda sheet: (normalize_header(sheet.title) != "feedbackexplorer", sheet.title))
    for sheet in candidates:
        for row_num in range(1, min(sheet.max_row, 15) + 1):
            headers = _header_map([cell.value for cell in sheet[row_num]])
            normalized = {normalize_header(item) for item in headers}
            has_id = bool(normalized.intersection(HEADER_ALIASES["feedback_id"]))
            has_text = bool(normalized.intersection(HEADER_ALIASES["text"]))
            if has_id and has_text:
                return sheet, row_num, headers
    raise ConversionError("No worksheet header containing a feedback ID and original text was found in the first 15 rows")


def _guess_model(path: Path) -> str | None:
    compact = normalize_header(path.stem)
    for model in ("F70", "G60", "G70"):
        if normalize_header(model) in compact:
            return model
    return None


def _feedback_model(value: Any, path: Path) -> tuple[str | None, str | None]:
    original = str(value).strip() if value is not None and str(value).strip() else None
    searchable = original or ""
    for code in ("F70", "G60", "G70", "G68"):
        if re.search(rf"(?i)(?<![A-Z0-9]){code}(?![A-Z0-9])", searchable):
            return code, original
    return _guess_model(path), original


def _normalize_region(region: Any, country: Any) -> tuple[str | None, str | None]:
    european_markets = {
        "albania", "andorra", "austria", "belarus", "belgium", "bosniaandherzegovina",
        "bulgaria", "croatia", "cyprus", "czechia", "czechrepublic", "denmark", "estonia",
        "finland", "france", "germany", "greece", "hungary", "iceland", "ireland", "italy",
        "kosovo", "latvia", "liechtenstein", "lithuania", "luxembourg", "malta", "moldova",
        "monaco", "montenegro", "netherlands", "northmacedonia", "poland", "portugal", "romania",
        "sanmarino", "serbia", "slovakia", "slovenia", "spain", "sweden", "switzerland",
        "ukraine", "vaticancity", "unitedkingdom", "gb", "uk", "norway", "no",
    }
    # Prefer an explicit commercial region, then fall back to the country.
    for raw_value in (region, country):
        raw = str(raw_value or "").strip()
        key = normalize_header(raw)
        if not key:
            continue
        if key in {"us", "usa", "unitedstates", "america"}:
            return "US", "commercial_region_map_v1"
        if key in {"cn", "china", "chinese"}:
            return "CN", "commercial_region_map_v1"
        if key in {"eu", "europe", "europeanunion", "european"} or key in european_markets:
            return "EU", "commercial_region_map_v1"
        if key in {"row", "restofworld", "othermarkets"}:
            return "RoW", "commercial_region_map_v1"
    # Country names are retained, but not silently assigned to a sales region.
    return None, None


def _normalize_powertrain(value: Any) -> tuple[str | None, str | None]:
    original = str(value).strip() if value is not None and str(value).strip() else None
    if original is None:
        return None, None
    normalized = "BEV" if normalize_header(original) == "beve" else original
    return original, normalized


def _text_quality(text: Any) -> str:
    normalized = re.sub(r"\s+", " ", str(text or "")).strip()
    if not normalized:
        return "empty"
    if normalized.casefold() in {"n/a", "na", "none", "null", "-", "?", "unknown", "not provided"}:
        return "low_information_candidate"
    if not re.search(r"[A-Za-zÀ-ÿ]{2,}", normalized):
        return "low_information_candidate"
    return "retained"


def convert_feedback_workbook(path: Path, source: Mapping[str, Any] | None = None) -> ConversionBundle:
    """Convert all feedback rows, preserving canonical comments and annotations."""
    bundle = ConversionBundle()
    source = dict(source or _source_file(path, "feedback", None))
    wb = _excel_module().load_workbook(path, read_only=True, data_only=False)
    try:
        sheet, header_row, headers = _find_feedback_sheet(wb)
        for context_row in range(1, header_row):
            context_cells = [cell.value for cell in sheet[context_row]]
            if not _nonempty(context_cells):
                continue
            context_headers = [f"column_{index}" for index in range(1, len(context_cells) + 1)]
            context_record = _row_source_record(source, sheet.title, context_row, context_headers, context_cells)
            context_record["record_kind"] = "table_context"
            bundle.source_records.append(context_record)
        header_cells = [cell.value for cell in sheet[header_row]]
        header_record = _row_source_record(source, sheet.title, header_row, headers, header_cells)
        header_record["record_kind"] = "table_header"
        bundle.source_records.append(header_record)
        staged: list[dict[str, Any]] = []
        # Read-only worksheets are forward-only XML streams. Repeated
        # ``sheet[row_num]`` lookups restart that stream for every row and
        # make large exports effectively quadratic. Consume the data once.
        data_rows = sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row,
                                    values_only=True)
        for row_num, values in enumerate(data_rows, start=header_row + 1):
            cells = list(values)
            if not _nonempty(cells):
                continue
            record = _row_source_record(source, sheet.title, row_num, headers, cells)
            record["original_record_id"] = json_value(_alias_value(headers, cells, "feedback_id"))
            bundle.source_records.append(record)
            staged.append({"row_num": row_num, "headers": headers, "cells": cells, "record": record})
        if not staged:
            raise ConversionError("Feedback worksheet has a valid header but no populated data rows")

        # Group first, so inconsistent reuse of a business ID is reported and
        # each conflicting row remains separately addressable.
        by_business_id: dict[str, list[dict[str, Any]]] = {}
        for item in staged:
            original = item["record"]["original_record_id"]
            if original is not None and str(original).strip():
                by_business_id.setdefault(str(original).strip(), []).append(item)

        signature_fields = ("text", "model", "country", "region", "powertrain", "source_label")
        signatures: dict[int, tuple[Any, ...]] = {}
        parsed_fields: dict[int, dict[str, Any]] = {}
        for item in staged:
            headers, cells = item["headers"], item["cells"]
            text = _alias_value(headers, cells, "text")
            country = _alias_value(headers, cells, "country")
            raw_region = _alias_value(headers, cells, "region")
            mapped_region, region_map = _normalize_region(raw_region, country)
            original_powertrain, powertrain = _normalize_powertrain(_alias_value(headers, cells, "powertrain"))
            row_model = _alias_value(headers, cells, "model")
            model, model_original = _feedback_model(row_model, path)
            fields = {
                "text": str(text or ""),
                "model": model,
                "model_original": model_original,
                "country": str(country).strip() if country is not None and str(country).strip() else None,
                "region": mapped_region,
                "region_original": str(raw_region).strip() if raw_region is not None and str(raw_region).strip() else None,
                "region_mapping_version": region_map,
                "powertrain_original": original_powertrain,
                "powertrain": powertrain,
                "source_label": str(_alias_value(headers, cells, "source_label")).strip()
                if _alias_value(headers, cells, "source_label") is not None else None,
                "feedback_type": str(_alias_value(headers, cells, "feedback_type")).strip()
                if _alias_value(headers, cells, "feedback_type") is not None else None,
                "vfc2": str(_alias_value(headers, cells, "vfc2")).strip()
                if _alias_value(headers, cells, "vfc2") is not None else None,
                "vfc3": str(_alias_value(headers, cells, "vfc3")).strip()
                if _alias_value(headers, cells, "vfc3") is not None else None,
            }
            parsed_fields[item["row_num"]] = fields
            signatures[item["row_num"]] = tuple(fields[name] for name in signature_fields)
            item["record"]["model"] = fields["model"]
            item["record"]["model_original"] = fields["model_original"]
            item["record"]["region"] = fields["region"]
            item["record"]["powertrain"] = fields["powertrain"]
            item["record"]["source_label"] = fields["source_label"]

        conflicts: set[str] = set()
        for business_id, items in by_business_id.items():
            if len({signatures[item["row_num"]] for item in items}) > 1:
                conflicts.add(business_id)
                bundle.issues.append({
                    "severity": "error",
                    "code": "duplicate_feedback_id_conflict",
                    "source_file_id": source["id"],
                    "business_id": business_id,
                    "rows": [item["row_num"] for item in items],
                    "message": "Repeated feedback ID has inconsistent text or scope metadata; rows were not merged.",
                })

        canonical_by_business_id: dict[str, str] = {}
        for business_id, items in by_business_id.items():
            if business_id not in conflicts:
                canonical_by_business_id[business_id] = stable_id("FB", source["source_version"], business_id)

        emitted_comments: set[str] = set()
        for item in staged:
            row_num = item["row_num"]
            record = item["record"]
            fields = parsed_fields[row_num]
            business_id = str(record["original_record_id"]).strip() if record["original_record_id"] is not None else ""
            if business_id and business_id not in conflicts:
                comment_id = canonical_by_business_id[business_id]
            else:
                comment_id = stable_id("FB", source["source_version"], "row", row_num, business_id)
            record["feedback_comment_id"] = comment_id
            if comment_id not in emitted_comments:
                bundle.feedback_comments.append({
                    "id": comment_id,
                    "source_file_id": source["id"],
                    "source_version": source["source_version"],
                    "original_id": business_id or None,
                    **fields,
                    "original_text": fields["text"],
                    "text_quality": _text_quality(fields["text"]),
                    "text_length": len(fields["text"]),
                    "canonicalization": "source_version_plus_original_id" if business_id and business_id not in conflicts else "source_row_fallback",
                })
                emitted_comments.add(comment_id)
            comment_text_quality = _text_quality(fields["text"])
            # ``text_quality`` belongs to the emitted comment payload, not to
            # the parsed workbook fields. Keep the warning test based on the
            # same computed value without requiring a nonexistent field.
            if comment_text_quality == "low_information_candidate":
                bundle.issues.append({
                    "severity": "warning",
                    "code": "low_information_text_candidate",
                    "source_file_id": source["id"],
                    "source_record_id": record["id"],
                    "message": "Source row is retained; downstream discovery may exclude it after review.",
                })

            if any(fields.get(name) for name in ("feedback_type", "vfc2", "vfc3")):
                annotation_id = stable_id("ANN", record["id"], 0)
                bundle.feedback_annotations.append({
                    "id": annotation_id,
                    "feedback_id": comment_id,
                    "source_record_id": record["id"],
                    "source_file_id": source["id"],
                    "feedback_type_original": fields["feedback_type"],
                    "vfc2_original": fields["vfc2"],
                    "vfc3_original": fields["vfc3"],
                    "assignment_method": "imported_annotation",
                    "assignment_review": "unreviewed",
                    "polarity": _polarity(fields["feedback_type"]),
                })
        bundle.source_files.append(dict(source))
        row_counts: dict[str, int] = {}
        comment_ids: dict[str, set[str]] = {}
        for item in staged:
            model = parsed_fields[item["row_num"]].get("model") or "unknown"
            row_counts[model] = row_counts.get(model, 0) + 1
            business_id = str(item["record"].get("original_record_id") or "").strip()
            if business_id and business_id not in conflicts:
                canonical_id = canonical_by_business_id[business_id]
            else:
                canonical_id = stable_id("FB", source["source_version"], "row", item["row_num"], business_id)
            comment_ids.setdefault(model, set()).add(canonical_id)
        bundle.manifest = {
            "source_file_id": source["id"],
            "kind": "feedback",
            "sheet": sheet.title,
            "header_row": header_row,
            "source_rows": len(staged),
            "canonical_comments": len(bundle.feedback_comments),
            "annotations": len(bundle.feedback_annotations),
            "duplicate_id_conflicts": len(conflicts),
            "model_counts": {model: {"source_rows": row_counts[model],
                                     "canonical_comments": len(comment_ids.get(model, set()))}
                             for model in sorted(row_counts)},
        }
        return bundle
    finally:
        wb.close()


def _polarity(feedback_type: Any) -> str | None:
    key = normalize_header(feedback_type)
    if key in {"like", "likes", "positive", "positivefeedback"}:
        return "positive"
    if key in {"difficulttouse", "dtu", "defect", "negative", "negativefeedback"}:
        return "negative"
    if key in {"wants", "request", "featurewish"}:
        return "request"
    return None


def _numeric(value: Any) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        numeric = float(value)
        if not math.isfinite(numeric):
            return None
        return int(numeric) if numeric.is_integer() else numeric
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text or not re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", text):
            return None
        number = float(text)
        return int(number) if number.is_integer() else number
    return None


def _year(value: Any) -> int | None:
    text = str(value or "").strip()
    # Underscores count as word characters, so ``volume_2030`` does not have
    # a ``\b`` before the year. Treat non-digits as valid delimiters.
    match = re.search(r"(?<!\d)(20(?:24|25|30))(?!\d)", text)
    return int(match.group(1)) if match else None


def _region(value: Any) -> str | None:
    mapped, _ = _normalize_region(value, value)
    return mapped


def _line(value: Any) -> str | None:
    compact = normalize_header(value)
    if "g60g68" in compact or ("g60" in compact and "g68" in compact):
        return "G60_G68"
    for code in ("F70", "G60", "G68", "G70"):
        if normalize_header(code) in compact:
            return code
    return None


def _scope_from_text(values: Iterable[Any]) -> tuple[str | None, str | None]:
    joined = " ".join(str(value) for value in values if value is not None)
    line = _line(joined)
    # In native study input, G68_CN remains derivative G68; only the sales
    # line parser maps the combined G60_G68 line.
    model = "G68" if re.search(r"(?i)\bG68\s*[_-]\s*CN\b", joined) else line
    region = None
    for candidate, aliases in {
        "US": ("United States", "USA", "US"),
        "CN": ("China", "CN"),
        "EU": ("Europe", "EU", "United Kingdom", "Norway", "GB"),
        "RoW": ("Rest of World", "RoW"),
    }.items():
        if any(re.search(rf"(?i)(?<![A-Za-z]){re.escape(alias)}(?![A-Za-z])", joined) for alias in aliases):
            region = candidate
            break
    return model, region


def _survey_column_scope(row: Mapping[str, Any], column_index: int) -> tuple[str | None, str | None]:
    """Resolve model/market from this column's headings, including merged cells."""
    header_values = row.get("header_values", [])
    column_header = header_values[column_index] if column_index < len(header_values) else None
    model = _line(column_header)
    model_headers = {candidate for value in header_values if (candidate := _line(value))}
    fallback_model, fallback_region = _scope_from_text([row.get("sheet")])
    if model is None and len(model_headers) <= 1:
        model = fallback_model or next(iter(model_headers), None)

    region = None
    for context_row in row.get("record", {}).get("raw_payload", {}).get("header_context", []):
        values = context_row.get("values", [])
        nearest_region = None
        # Merged spreadsheet headings leave continuation cells blank. Carry a
        # region label across those columns, stopping at the next explicit one.
        for index in range(min(column_index + 1, len(values))):
            candidate = _region(values[index])
            if candidate:
                nearest_region = candidate
        if nearest_region:
            region = nearest_region
    if region is None:
        region = fallback_region
    return model, region


def _headers_and_records(path: Path, source: Mapping[str, Any], bundle: ConversionBundle) -> list[dict[str, Any]]:
    openpyxl = _excel_module()
    # Keep formula text in the source record while also reading Excel's cached
    # result for numeric facts. openpyxl does not calculate formulas itself.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(wb.worksheets):
            cached_sheet = cached_wb.worksheets[sheet_index]
            width = sheet.max_column
            # A title/merged heading often precedes the table. Prefer the
            # densest of the first 20 populated rows, while retaining every
            # populated row with its original locator below.
            populated = [(n, [cell.value for cell in sheet[n]])
                         for n in range(1, min(sheet.max_row, 20) + 1)
                         if _nonempty(cell.value for cell in sheet[n])]
            def header_score(values: Sequence[Any]) -> tuple[float, int, int]:
                nonempty_count = sum(1 for value in values if value is not None and str(value).strip())
                score = 0.0
                for value in values:
                    if value is None:
                        continue
                    normalized = normalize_header(value)
                    if any(normalized in aliases for aliases in HEADER_ALIASES.values()):
                        score += 4
                    # Product/derivative headings are the column schema in
                    # BMW's wide study sheets, even when no literal "model"
                    # header exists.
                    if _line(value):
                        score += 6
                    if _year(value) is not None or _region(value) is not None:
                        score += 2
                    if normalized in {str(number) for number in range(1, 8)}:
                        score += 1
                    if any(label in normalized for label in
                           ("satisfied", "dissatisfied", "agree", "disagree", "important", "quality", "rating")):
                        score += 1
                return score, nonempty_count, -len(values)

            header_row = max(populated, key=lambda item: header_score(item[1]))[0] if populated else 1
            headers = _header_map([cell.value for cell in sheet[header_row]])
            header_values = [cell.value for cell in sheet[header_row]]
            header_context = [
                {"row": row_num, "values": [json_value(value) for value in values]}
                for row_num, values in populated if row_num <= header_row
            ]
            question_label = (str(header_values[0]).strip()
                              if header_values and header_values[0] is not None else None)
            # Iterate the read-only worksheet once instead of calling
            # ``sheet.cell`` for every cell, which repeatedly scans the XML.
            formula_rows = sheet.iter_rows(min_row=1, max_row=sheet.max_row,
                                           min_col=1, max_col=width, values_only=True)
            cached_rows = cached_sheet.iter_rows(min_row=1, max_row=sheet.max_row,
                                                 min_col=1, max_col=width, values_only=True)
            for row_num, (values, cached_values) in enumerate(
                    zip(formula_rows, cached_rows, strict=True), start=1):
                cells = list(values)
                cached_cells = list(cached_values)
                if not _nonempty(cells):
                    continue
                source_record = _row_source_record(source, sheet.title, row_num, headers, cells)
                source_record["raw_payload"]["header_context"] = header_context
                source_record["record_kind"] = "table_header" if row_num == header_row else "original_row"
                bundle.source_records.append(source_record)
                if row_num <= header_row:
                    continue
                first_value = cells[0] if cells else None
                has_numeric = any(
                    _effective_numeric(value, cached_cells[index] if index < len(cached_cells) else None)[0] is not None
                    for index, value in enumerate(cells)
                )
                measure_label = (str(first_value).strip()
                                 if has_numeric and first_value is not None and str(first_value).strip()
                                 else None)
                if not has_numeric and first_value is not None and str(first_value).strip():
                    question_label = str(first_value).strip()
                rows.append({"sheet": sheet.title, "row_num": row_num, "headers": headers,
                             "header_values": header_values, "cells": cells,
                             "cached_cells": cached_cells, "question_label": question_label,
                             "measure_label": measure_label, "record": source_record})
    finally:
        wb.close()
        cached_wb.close()
    return rows


def _effective_numeric(raw_value: Any, cached_value: Any) -> tuple[int | float | None, str]:
    """Return a numeric source value and whether it came from a formula cache."""
    if isinstance(raw_value, str) and raw_value.startswith("="):
        return _numeric(cached_value), "cached_formula_result"
    return _numeric(raw_value), "source_numeric_cell"


def _fact(source: Mapping[str, Any], record: Mapping[str, Any], kind: str,
          locator: Mapping[str, Any], **properties: Any) -> dict[str, Any]:
    fact_id = stable_id("FACT", source["id"], kind, json.dumps(locator, sort_keys=True), properties.get("metric"), properties.get("value"))
    return {
        "id": fact_id,
        "source_file_id": source["id"],
        "source_record_id": record["id"],
        "source_version": source["source_version"],
        "kind": kind,
        "locator": dict(locator),
        **properties,
    }


def _convert_survey_rows(rows: Sequence[Mapping[str, Any]], source: Mapping[str, Any], bundle: ConversionBundle) -> None:
    for row in rows:
        headers, cells = row["headers"], row["cells"]
        cached_cells = row.get("cached_cells", cells)
        respondent_n_value = _numeric(_alias_value(headers, cells, "respondent_n"))
        respondent_n = int(respondent_n_value) if respondent_n_value is not None and respondent_n_value >= 0 else None
        raw_scale = _alias_value(headers, cells, "scale")
        scale = str(raw_scale).strip() if raw_scale is not None and str(raw_scale).strip() else None
        raw_unit = _alias_value(headers, cells, "unit")
        # Some BMW sheets put a measure label such as "% of column, sample
        # base" in the second cell rather than a named column.
        if raw_unit is None and len(cells) > 1 and isinstance(cells[1], str):
            raw_unit = cells[1]
        unit = str(raw_unit).strip() if raw_unit is not None and str(raw_unit).strip() else None
        question = row.get("question_label")
        measure_label = row.get("measure_label")
        numeric_cells = []
        for index, raw_value in enumerate(cells):
            cached_value = cached_cells[index] if index < len(cached_cells) else None
            value, value_provenance = _effective_numeric(raw_value, cached_value)
            if value is not None:
                numeric_cells.append((index, value, value_provenance))
        if not numeric_cells:
            continue
        # Preserve every numeric cell as a native fact, including scales and
        # percentages. No cross-scale averaging or respondent-N inference.
        for index, value, value_provenance in numeric_cells:
            column_header = headers[index] if index < len(headers) else f"column_{index + 1}"
            model, region = _survey_column_scope(row, index)
            metric_parts = [part for part in (question, measure_label) if part]
            metric = " / ".join(dict.fromkeys(metric_parts)) or str(column_header)
            cell_locator = {**row["record"]["locator"], "column": index + 1,
                            "column_header": column_header}
            bundle.facts.append(_fact(
                source, row["record"], "survey_cell", cell_locator,
                metric=metric, question=question, response_label=measure_label,
                value=value, unit=unit, scale=scale,
                derivative=model, region=region, respondent_n=respondent_n,
                interpretation="native_cell_unmapped", mapping_strength="unmapped",
                source_value=json_value(cells[index]), value_provenance=value_provenance,
            ))

        # Recognize a seven-column 1..7 distribution only when those native
        # category labels are explicit in the workbook header.
        category_positions: list[int] = []
        for index, header in enumerate(row["header_values"]):
            if str(header).strip() in {"1", "2", "3", "4", "5", "6", "7"}:
                category_positions.append(index)
        if len(category_positions) == 7 and category_positions == list(range(category_positions[0], category_positions[0] + 7)):
            values = [
                _effective_numeric(
                    cells[index] if index < len(cells) else None,
                    cached_cells[index] if index < len(cached_cells) else None,
                )[0]
                for index in category_positions
            ]
            if all(value is not None for value in values):
                question_labels = [str(value).strip() for value in cells[:category_positions[0]]
                                   if isinstance(value, str) and str(value).strip()]
                question = " / ".join(question_labels) or row["sheet"]
                distribution = {str(row["header_values"][index]): value
                                for index, value in zip(category_positions, values)}
                locators = [{**row["record"]["locator"], "column": index + 1,
                             "column_header": str(row["header_values"][index])}
                            for index in category_positions]
                bundle.facts.append(_fact(
                    source, row["record"], "survey_distribution",
                    {**row["record"]["locator"], "cells": locators},
                    metric=question, question=question, distribution=distribution,
                    distribution_total=sum(values), distribution_kind="native_unclassified", unit=unit,
                    category_labels=[str(row["header_values"][index]) for index in category_positions],
                    respondent_n=respondent_n,
                    scale=scale or "1-7 (explicit native category headers)",
                    scale_provenance="explicit_source_column" if scale else "derived_from_explicit_category_headers",
                    derivative=model, region=region,
                    mapping_strength="unmapped",
                    note="Seven native categories preserved; total is not interpreted as respondent N.",
                ))

    _add_native_survey_distributions(rows, source, bundle)


def _add_native_survey_distributions(rows: Sequence[Mapping[str, Any]],
                                     source: Mapping[str, Any], bundle: ConversionBundle) -> None:
    """Bundle seven source response rows without inventing N or a mapped topic."""
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for row in rows:
        category = row.get("measure_label")
        cells = row["cells"]
        cached_cells = row.get("cached_cells", cells)
        if not category or normalize_header(category) in {"sampletotal", "mean"}:
            continue
        source_measure = cells[1] if len(cells) > 1 and isinstance(cells[1], str) else ""
        if "%" not in source_measure and "percent" not in source_measure.casefold():
            continue
        for index, raw_value in enumerate(cells):
            cached_value = cached_cells[index] if index < len(cached_cells) else None
            value, value_provenance = _effective_numeric(raw_value, cached_value)
            if value is None:
                continue
            model, region = _survey_column_scope(row, index)
            if not model or not region:
                continue
            key = (str(row["sheet"]), str(row.get("question_label") or ""), model, region)
            grouped.setdefault(key, []).append({
                "category": str(category), "value": value, "value_provenance": value_provenance,
                "row": row, "column": index,
            })

    for (sheet, question, model, region), entries in grouped.items():
        categories = [entry["category"] for entry in entries]
        if len(entries) != 7 or len(set(categories)) != 7:
            continue
        total = sum(entry["value"] for entry in entries)
        if not (abs(total - 1.0) <= 0.03 or abs(total - 100.0) <= 3.0):
            continue
        ordered = sorted(entries, key=lambda entry: entry["row"]["row_num"])
        first = ordered[0]
        locators = []
        for entry in ordered:
            entry_row = entry["row"]
            column_index = entry["column"]
            column_header = (entry_row["header_values"][column_index]
                             if column_index < len(entry_row["header_values"]) else None)
            locators.append({
                **entry_row["record"]["locator"],
                "source_record_id": entry_row["record"]["id"],
                "column": column_index + 1,
                "column_header": json_value(column_header),
                "response_category": entry["category"],
                "value_provenance": entry["value_provenance"],
            })
        unit = first["row"]["cells"][1] if len(first["row"]["cells"]) > 1 else None
        bundle.facts.append(_fact(
            source, first["row"]["record"], "survey_distribution",
            {"filename": source["filename"], "sheet": sheet, "question": question, "cells": locators},
            metric=question, question=question,
            distribution={entry["category"]: entry["value"] for entry in ordered},
            distribution_total=total, distribution_kind="native_seven_category",
            category_labels=[entry["category"] for entry in ordered],
            respondent_n=None, scale=None, scale_provenance="not_reported",
            unit=unit, derivative=model, region=region,
            mapping_strength="unmapped",
            note="Native seven-category shares; sample base is not interpreted as respondent N.",
            cell_locators=locators,
        ))


def _convert_sales_rows(rows: Sequence[Mapping[str, Any]], source: Mapping[str, Any], bundle: ConversionBundle) -> None:
    for row in rows:
        headers, cells = row["headers"], row["cells"]
        cached_cells = row.get("cached_cells", cells)
        header_context = row["record"].get("raw_payload", {}).get("header_context", [])
        context_values = [value for context_row in header_context for value in context_row.get("values", [])]
        context = [row["sheet"], source["filename"], *cells, *row["header_values"], *context_values]
        context_years = {_year(value) for value in context_values if _year(value)}
        context_regions = {_region(value) for value in context_values if _region(value)}
        line = _line(" ".join(str(value) for value in context if value is not None))
        # Prefer explicit long-table columns if present.
        explicit_year = _alias_value(headers, cells, "year")
        explicit_volume = _alias_value(headers, cells, "volume")
        explicit_region = _alias_value(headers, cells, "region")
        explicit_line = _alias_value(headers, cells, "model")
        if explicit_line:
            line = _line(explicit_line) or line
        if explicit_year is not None and explicit_volume is not None and explicit_region is not None:
            year = _year(explicit_year)
            region = _region(explicit_region)
            volume_index = next((index for index, header in enumerate(headers)
                                 if normalize_header(header) in HEADER_ALIASES["volume"]), None)
            cached_volume = cached_cells[volume_index] if volume_index is not None and volume_index < len(cached_cells) else None
            volume, value_provenance = _effective_numeric(explicit_volume, cached_volume)
            if year and region and line and volume is not None:
                bundle.facts.append(_fact(
                    source, row["record"], "sales", row["record"]["locator"],
                    vehicle_line=line, derivative=None, region=region, year=year,
                    volume=volume, unit="vehicles", powertrain=None,
                    powertrain_split="unknown", applicability="whole_line",
                    mapping_version="sales_line_map_v1",
                    source_value=json_value(explicit_volume), value_provenance=value_provenance,
                ))
                continue

        # Matrix layout: year/region may be on either the row or column header.
        for index, raw_value in enumerate(cells):
            cached_value = cached_cells[index] if index < len(cached_cells) else None
            volume, value_provenance = _effective_numeric(raw_value, cached_value)
            if isinstance(raw_value, str) and raw_value.startswith("=") and volume is None:
                bundle.issues.append({
                    "severity": "warning", "code": "sales_formula_without_cached_value",
                    "source_file_id": source["id"], "source_record_id": row["record"]["id"],
                    "column": index + 1,
                    "message": "Formula text was preserved, but the workbook has no cached numeric result to use as a sales fact.",
                })
            if volume is None:
                continue
            column_label = row["header_values"][index] if index < len(row["header_values"]) else None
            # In matrix layouts an exact 2024/2025/2030 cell can be a row
            # label. Keep it in the raw row and use it as context for sibling
            # volume cells, rather than importing it as a vehicle volume.
            if (_year(raw_value) is not None and volume == _year(raw_value)
                    and _year(column_label) is None):
                continue
            row_values = [value for value in cells if value is not None]
            year = (_year(column_label) or next((_year(value) for value in row_values if _year(value)), None)
                    or (next(iter(context_years)) if len(context_years) == 1 else None))
            region = (_region(column_label) or next((_region(value) for value in row_values if _region(value)), None)
                      or (next(iter(context_regions)) if len(context_regions) == 1 else None))
            if not (line and year and region):
                continue
            locator = {**row["record"]["locator"], "column": index + 1,
                       "column_header": json_value(column_label)}
            bundle.facts.append(_fact(
                source, row["record"], "sales", locator,
                vehicle_line=line, derivative=None, region=region, year=year,
                volume=volume, unit="vehicles", powertrain=None,
                powertrain_split="unknown", applicability="whole_line",
                mapping_version="sales_line_map_v1",
                source_value=json_value(raw_value), value_provenance=value_provenance,
            ))


def convert_workbook(path: Path, kind: str, source: Mapping[str, Any] | None = None) -> ConversionBundle:
    """Convert study or sales workbook while retaining every populated row."""
    source = dict(source or _source_file(path, kind, None))
    bundle = ConversionBundle(source_files=[source])
    rows = _headers_and_records(path, source, bundle)
    if not rows:
        raise ConversionError("Workbook contains no populated source rows")
    if kind == "survey":
        _convert_survey_rows(rows, source, bundle)
        if not bundle.facts:
            bundle.issues.append({"severity": "error", "code": "survey_no_numeric_facts",
                                  "source_file_id": source["id"],
                                  "message": "Workbook rows were preserved, but no numeric study cells could be converted."})
    elif kind == "sales":
        _convert_sales_rows(rows, source, bundle)
        if not any(fact.get("kind") == "sales" for fact in bundle.facts):
            bundle.issues.append({"severity": "error", "code": "sales_no_scoped_facts",
                                  "source_file_id": source["id"],
                                  "message": "Workbook rows were preserved, but no sales value had an identifiable line, region, and year."})
    else:
        for row in rows:
            for index, value in enumerate(row["cells"]):
                if _numeric(value) is None:
                    continue
                locator = {**row["record"]["locator"], "column": index + 1}
                bundle.facts.append(_fact(source, row["record"], "workbook_numeric_cell", locator,
                                          metric=row["headers"][index] if index < len(row["headers"]) else None,
                                          value=_numeric(value), interpretation="unmapped"))
    bundle.manifest = {
        "source_file_id": source["id"], "kind": kind,
        "source_rows": len(rows), "facts": len(bundle.facts),
        "unmapped_numeric_cells": sum(1 for fact in bundle.facts if fact.get("mapping_strength") == "unmapped"),
    }
    return bundle


def _chunk_page_text(text: str, max_words: int = 850) -> list[tuple[int, int, str]]:
    spans = list(re.finditer(r"\S+", text))
    chunks: list[tuple[int, int, str]] = []
    for offset in range(0, len(spans), max_words):
        group = spans[offset:offset + max_words]
        if not group:
            continue
        start, end = group[0].start(), group[-1].end()
        chunks.append((start, end, text[start:end]))
    return chunks


def convert_pdf(path: Path, source: Mapping[str, Any] | None = None,
                document_kind: str | None = None) -> ConversionBundle:
    """Extract page-preserving PDF text and exact per-page chunk substrings."""
    document_kind = document_kind or classify_source(path)[1] or "document"
    source = dict(source or _source_file(path, "document", document_kind))
    bundle = ConversionBundle(source_files=[source])
    reader = _pdf_reader(path)
    if len(reader.pages) == 0:
        raise ConversionError("PDF contains no pages")
    page_summaries: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # preserve per-page failures, continue safely
            text = ""
            bundle.issues.append({"severity": "error", "code": "pdf_page_extraction_failed",
                                  "source_file_id": source["id"], "page": page_number,
                                  "message": f"{type(exc).__name__}: {exc}"})
        record_id = stable_id("SR", source["id"], "page", page_number)
        record = {
            "id": record_id, "source_file_id": source["id"],
            "source_version": source["source_version"],
            "locator": {"filename": source["filename"], "page": page_number, "printed_page": None},
            "raw_payload": {"text": text, "extraction": "pypdf", "page_text_empty": not bool(text.strip())},
            "original_record_id": None,
        }
        bundle.source_records.append(record)
        chunk_count = 0
        for chunk_number, (start, end, chunk_text) in enumerate(_chunk_page_text(text), start=1):
            chunk_id = stable_id("DOC", source["id"], page_number, chunk_number, start, end)
            bundle.document_chunks.append({
                "id": chunk_id, "source_file_id": source["id"],
                "source_record_id": record_id, "source_version": source["source_version"],
                "document_kind": document_kind, "page": page_number, "printed_page": None,
                "chunk_number": chunk_number, "start_offset": start, "end_offset": end,
                "original_text": chunk_text, "derivative": None, "market": None,
                "applicability": "unknown",
            })
            chunk_count += 1
        if not text.strip():
            bundle.issues.append({"severity": "error", "code": "pdf_page_empty_or_unextractable",
                                  "source_file_id": source["id"], "page": page_number,
                                  "message": "Page retained with an empty text record; citation text is unavailable."})
        page_summaries.append({"page": page_number, "characters": len(text), "chunks": chunk_count})
    bundle.manifest = {"source_file_id": source["id"], "kind": "document",
                       "document_kind": document_kind, "pages": page_summaries,
                       "page_count": len(page_summaries), "chunks": len(bundle.document_chunks)}
    return bundle


def _merge_bundle(target: ConversionBundle, source: ConversionBundle) -> None:
    for field_name in ("source_files", "source_records", "feedback_comments", "feedback_annotations",
                       "facts", "document_chunks", "issues"):
        getattr(target, field_name).extend(getattr(source, field_name))
    target.manifest.setdefault("sources", []).append(source.manifest)


def discover_sources(raw_data_dir: Path) -> list[Path]:
    if not raw_data_dir.exists() or not raw_data_dir.is_dir():
        raise ConversionError(f"Raw data directory does not exist or is not a directory: {raw_data_dir}")
    return sorted((path for path in raw_data_dir.rglob("*") if path.is_file()
                   and not path.name.startswith("~$")
                   and path.suffix.casefold() in SUPPORTED_EXCEL | SUPPORTED_PDF),
                  key=lambda item: str(item.relative_to(raw_data_dir)).casefold())


EXPECTED_FEEDBACK_COUNTS = {
    "F70": {"source_rows": 4076, "canonical_comments": 2128},
    "G60": {"source_rows": 5005, "canonical_comments": 4000},
    "G70": {"source_rows": 4690, "canonical_comments": 4000},
}


def convert_sources(raw_data_dir: Path, required_kinds: Sequence[str] = DEFAULT_REQUIRED_KINDS,
                    check_known_feedback_counts: bool = False,
                    required_document_kinds: Sequence[str] = DEFAULT_REQUIRED_DOCUMENT_KINDS) -> ConversionBundle:
    """Convert supported files without writing to, renaming, or deleting inputs."""
    raw_data_dir = raw_data_dir.resolve()
    paths = discover_sources(raw_data_dir)
    if not paths:
        raise ConversionError("No supported XLSX/XLSM/PDF inputs found; no empty success bundle was produced")
    bundle = ConversionBundle()
    source_kinds: set[str] = set()
    seen_source_ids: dict[str, str] = {}
    for path in paths:
        kind, document_kind = classify_source(path)
        try:
            source = _source_file(path, kind, document_kind)
            prior_kind = seen_source_ids.get(source["id"])
            if prior_kind is not None:
                severity = "warning" if prior_kind == kind else "error"
                bundle.issues.append({"severity": severity,
                                      "code": "duplicate_source_content_skipped",
                                      "filename": path.name, "source_file_id": source["id"],
                                      "kind": kind, "prior_kind": prior_kind,
                                      "message": "Identical file content is imported once per SHA-256; duplicate path skipped."})
                continue
            seen_source_ids[source["id"]] = kind
            if kind == "feedback":
                converted = convert_feedback_workbook(path, source)
            elif kind in {"survey", "sales", "workbook"}:
                converted = convert_workbook(path, kind, source)
            elif kind == "document":
                converted = convert_pdf(path, source, document_kind)
            else:
                bundle.issues.append({"severity": "warning", "code": "unsupported_source",
                                      "filename": path.name, "message": "Input extension is not supported."})
                continue
            source_kinds.add(kind)
            _merge_bundle(bundle, converted)
        except Exception as exc:
            bundle.issues.append({"severity": "error", "code": "source_conversion_failed",
                                  "filename": path.name, "kind": kind,
                                  "message": f"{type(exc).__name__}: {exc}"})

    missing_kinds = sorted(set(required_kinds) - source_kinds)
    for kind in missing_kinds:
        bundle.issues.append({"severity": "error", "code": "required_input_kind_missing",
                              "kind": kind, "message": f"No successfully converted source of kind '{kind}' was found."})
    available_document_kinds = sorted({
        str(source.get("document_kind")) for source in bundle.source_files
        if source.get("kind") == "document" and source.get("document_kind")
    })
    missing_document_kinds = sorted(set(required_document_kinds) - set(available_document_kinds))
    for document_kind in missing_document_kinds:
        bundle.issues.append({
            "severity": "error", "code": "required_document_kind_missing",
            "document_kind": document_kind,
            "message": f"No successfully converted document of kind '{document_kind}' was found.",
        })
    successful_source_count = len(bundle.source_files)
    observed_feedback: dict[str, dict[str, int]] = {}
    for source_summary in bundle.manifest.get("sources", []):
        if source_summary.get("kind") != "feedback":
            continue
        for model, model_counts in source_summary.get("model_counts", {}).items():
            observed = observed_feedback.setdefault(model, {"source_rows": 0, "canonical_comments": 0})
            observed["source_rows"] += int(model_counts.get("source_rows", 0))
            observed["canonical_comments"] += int(model_counts.get("canonical_comments", 0))
    known_count_check = {}
    for model, expected in EXPECTED_FEEDBACK_COUNTS.items():
        observed = observed_feedback.get(model)
        known_count_check[model] = {
            "expected": expected,
            "observed": observed,
            "status": "not_present" if observed is None else ("matches" if observed == expected else "mismatch"),
        }
        if check_known_feedback_counts and observed is not None and observed != expected:
            bundle.issues.append({"severity": "error", "code": "feedback_count_acceptance_mismatch",
                                  "model": model, "expected": expected, "observed": observed,
                                  "message": "Observed counts differ from the documented complete-export target."})
    if check_known_feedback_counts:
        for model in EXPECTED_FEEDBACK_COUNTS:
            if model not in observed_feedback:
                known_count_check[model]["status"] = "not_present"
                bundle.issues.append({"severity": "error", "code": "feedback_count_acceptance_model_missing",
                                      "model": model,
                                      "message": "Complete-export acceptance check requested but no feedback source for this model was found."})
    if check_known_feedback_counts and any(issue.get("severity") == "error" for issue in bundle.issues):
        status = "partial" if successful_source_count else "failed"
    if successful_source_count == 0:
        status = "failed"
    elif missing_kinds or missing_document_kinds or any(issue["severity"] == "error" for issue in bundle.issues):
        status = "partial"
    else:
        status = "succeeded"
    bundle.manifest.update({
        "conversion_version": CONVERSION_VERSION,
        "status": status,
        "required_kinds": list(required_kinds),
        "required_document_kinds": list(required_document_kinds),
        "available_kinds": sorted(source_kinds),
        "available_document_kinds": available_document_kinds,
        "missing_kinds": missing_kinds,
        "missing_document_kinds": missing_document_kinds,
        "known_feedback_count_check_enabled": check_known_feedback_counts,
        "known_feedback_counts": known_count_check,
        "sources": bundle.manifest.get("sources", []),
        "counts": {
            "source_files": len(bundle.source_files),
            "source_records": len(bundle.source_records),
            "feedback_comments": len(bundle.feedback_comments),
            "feedback_annotations": len(bundle.feedback_annotations),
            "facts": len(bundle.facts),
            "document_chunks": len(bundle.document_chunks),
            "issues": len(bundle.issues),
        },
        "issues": bundle.issues,
    })
    return bundle


def _write_jsonl(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=json_value))
            stream.write("\n")


def write_bundle(bundle: ConversionBundle, output_dir: Path, raw_data_dir: Path) -> Path:
    """Write a versioned intermediate bundle outside the raw source directory."""
    output_dir = output_dir.resolve()
    raw_data_dir = raw_data_dir.resolve()
    if output_dir == raw_data_dir or raw_data_dir in output_dir.parents:
        raise ConversionError("Output directory must be outside the raw data directory; originals are immutable")
    output_dir.mkdir(parents=True, exist_ok=True)
    tables = {
        "source_files.jsonl": bundle.source_files,
        "source_records.jsonl": bundle.source_records,
        "feedback_comments.jsonl": bundle.feedback_comments,
        "feedback_annotations.jsonl": bundle.feedback_annotations,
        "facts.jsonl": bundle.facts,
        "document_chunks.jsonl": bundle.document_chunks,
    }
    for filename, rows in tables.items():
        _write_jsonl(output_dir / filename, rows)
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(bundle.manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return manifest_path
