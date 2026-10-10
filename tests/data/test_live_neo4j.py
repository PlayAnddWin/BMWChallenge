"""Opt-in end-to-end checks against a disposable, loopback Neo4j instance."""

from __future__ import annotations

import os
from pathlib import Path
import re
import tempfile
import unittest
from urllib.parse import urlparse
from uuid import uuid4

from openpyxl import Workbook

from config import get_settings
from services.evidence import EvidenceScope, Neo4jEvidenceProvider
from services.ingestion import convert_sources, write_bundle


REPO_ROOT = Path(__file__).resolve().parents[2]
REAL_DATA_DIR = get_settings().raw_data_dir
NEO4J_URI = os.environ.get("BMW_NEO4J_TEST_URI")


class LiveNeo4jTests(unittest.TestCase):
    """Import the real source set and prove scoped evidence reads with Neo4j."""

    @classmethod
    def setUpClass(cls) -> None:
        if not NEO4J_URI:
            raise unittest.SkipTest("set BMW_NEO4J_TEST_URI to run live Neo4j integration checks")
        host = urlparse(NEO4J_URI).hostname
        if host not in {"localhost", "127.0.0.1", "::1"}:
            raise RuntimeError("Live integration tests only write to a loopback Neo4j instance")
        try:
            from neo4j import GraphDatabase
        except ImportError as exc:
            raise RuntimeError("Install the Neo4j Python driver from pipeline/requirements-data.txt") from exc

        if os.environ.get("BMW_NEO4J_TEST_AUTH_DISABLED") == "1":
            auth = None
        else:
            settings = get_settings()
            user = os.environ.get("BMW_NEO4J_TEST_USER", settings.neo4j_user)
            password = os.environ.get("BMW_NEO4J_TEST_PASSWORD", settings.neo4j_password.get_secret_value())
            if not user or not password:
                raise RuntimeError(
                    "Set BMW_NEO4J_TEST_USER and BMW_NEO4J_TEST_PASSWORD, or use a loopback test DB "
                    "with BMW_NEO4J_TEST_AUTH_DISABLED=1"
                )
            auth = (user, password)
        cls.driver = GraphDatabase.driver(NEO4J_URI, auth=auth)
        cls.driver.verify_connectivity()

    @classmethod
    def tearDownClass(cls) -> None:
        if getattr(cls, "driver", None) is not None:
            cls.driver.close()

    def _loader(self):
        from importlib.util import spec_from_file_location, module_from_spec

        spec = spec_from_file_location("bmw_data_loader", REPO_ROOT / "pipeline" / "01_load_neo4j.py")
        assert spec is not None and spec.loader is not None
        loader = module_from_spec(spec)
        spec.loader.exec_module(loader)
        return loader

    def _graph_snapshot(self) -> dict:
        with self.driver.session() as session:
            nodes = session.run(
                "MATCH (n) RETURN labels(n) AS labels, count(n) AS count ORDER BY labels"
            ).data()
            relationships = session.run(
                "MATCH ()-[r]->() RETURN type(r) AS type, count(r) AS count ORDER BY type"
            ).data()
        return {"nodes": nodes, "relationships": relationships}

    def test_import_and_audit_failure_roll_back_new_evidence(self) -> None:
        token = uuid4().hex
        with tempfile.TemporaryDirectory(dir=REPO_ROOT) as temp_dir:
            root = Path(temp_dir)
            raw_dir = root / "raw"
            raw_dir.mkdir()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Feedback_Explorer"
            sheet.append(("ID", "Customer Feedback", "Model", "Country", "Feedback Type", "VFC2"))
            sheet.append((token, f"Synthetic rollback probe {token}.", "G70", "US", "DTU", token))
            workbook.save(raw_dir / "G70_feedback.xlsx")
            bundle = convert_sources(raw_dir, required_kinds=("feedback",), required_document_kinds=())
            self.assertEqual(bundle.manifest["status"], "succeeded")
            processed_dir = root / "processed"
            write_bundle(bundle, processed_dir, raw_dir)
            probe_ids = [row["id"] for table in (
                bundle.source_files, bundle.source_records, bundle.feedback_comments, bundle.feedback_annotations
            ) for row in table] + [token]

            def probe_count() -> int:
                with self.driver.session() as session:
                    return session.run(
                        "MATCH (n) WHERE n.id IN $ids OR n.vfc2 = $topic RETURN count(n) AS count",
                        ids=probe_ids, topic=token,
                    ).single()["count"]

            self.assertEqual(probe_count(), 0)

            def failing_audit_sink(tx, event) -> None:
                staged_count = tx.run(
                    "MATCH (sf:SourceFile) WHERE sf.id IN $ids RETURN count(sf) AS count",
                    ids=event["source_file_ids"],
                ).single()["count"]
                self.assertEqual(staged_count, 1)
                tx.run("CREATE (:Role1RollbackProbe {id: $id})", id=token).consume()
                raise RuntimeError("injected audit failure")

            with self.assertRaisesRegex(RuntimeError, "injected audit failure"):
                self._loader().load_processed_bundle(self.driver, processed_dir, audit_sink=failing_audit_sink)

            self.assertEqual(probe_count(), 0, "Evidence, taxonomy and audit writes must roll back together")

    def test_real_import_is_idempotent_and_citations_resolve(self) -> None:
        self.assertTrue(REAL_DATA_DIR.is_dir(), f"BMW data directory not found: {REAL_DATA_DIR}")
        bundle = convert_sources(REAL_DATA_DIR, check_known_feedback_counts=True)
        self.assertEqual(bundle.manifest["status"], "succeeded", bundle.manifest.get("issues"))
        self.assertIn("corporate_report", {
            source.get("document_kind") for source in bundle.source_files
        })

        loader = self._loader()

        scratch_dir = REPO_ROOT / "scratch"
        scratch_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="neo4j-integration-", dir=scratch_dir) as temp_dir:
            processed_dir = Path(temp_dir) / "processed"
            write_bundle(bundle, processed_dir, REAL_DATA_DIR)
            first = loader.load_processed_bundle(self.driver, processed_dir)
            first_snapshot = self._graph_snapshot()
            second = loader.load_processed_bundle(self.driver, processed_dir)
            second_snapshot = self._graph_snapshot()

        self.assertEqual(first, second)
        self.assertEqual(first_snapshot, second_snapshot, "Reimport must not grow any node or relationship count")
        self.assertEqual(first["counts"]["feedback_comments"], 10_128)
        self.assertEqual(first["counts"]["feedback_annotations"], 13_371)
        self.assertGreater(first["counts"]["document_chunks"], 78)

        with self.driver.session() as session:
            graph_counts = session.run("""
                MATCH (feedback:CustomerFeedback)
                WITH count(feedback) AS feedback_count
                MATCH (chunk:DocumentChunk {document_kind: 'corporate_report'})
                WITH feedback_count, count(chunk) AS strategy_chunk_count
                RETURN feedback_count, strategy_chunk_count
            """).single()
        self.assertEqual(graph_counts["feedback_count"], 10_128)
        self.assertGreater(graph_counts["strategy_chunk_count"], 0)
        with self.driver.session() as session:
            current_import_runs = session.run(
                "MATCH (run:ImportRun {id: $id}) RETURN count(run) AS count",
                id=first["import_run_id"],
            ).single()["count"]
        self.assertEqual(current_import_runs, 1)

        provider = Neo4jEvidenceProvider(self.driver)
        sample = next(comment for comment in bundle.feedback_comments
                      if comment.get("model") == "G70" and comment.get("region") == "US"
                      and comment.get("original_text", "").strip())
        searchable_words = re.findall(r"[A-Za-z0-9_'-]{6,}", sample["original_text"])
        self.assertTrue(searchable_words, "Chosen feedback sample has no searchable word")
        g70_us_texts = [comment.get("original_text", "").casefold()
                        for comment in bundle.feedback_comments
                        if comment.get("model") == "G70" and comment.get("region") == "US"]
        query = min(searchable_words,
                    key=lambda word: sum(word.casefold() in text for text in g70_us_texts))
        matches = provider.search_feedback(query, EvidenceScope(model="G70", region="US"), limit=100)
        sample_match = next((match for match in matches if match["id"] == sample["id"]), None)
        self.assertIsNotNone(sample_match)
        self.assertEqual(sample_match["model_original"], "G70")
        self.assertTrue(any(annotation.get("vfc2_original") for annotation in sample_match["annotations"]))
        source_record_id = sample_match["source_records"][0]["id"]
        source_record = provider.get_source_record(source_record_id)
        self.assertIsNotNone(source_record)
        self.assertTrue(source_record["locator"].get("row"))
        self.assertIn(sample["original_text"], source_record["raw_payload"]["values_by_header"].values())

        report_chunks = provider.search_document_chunks(
            "BMW", document_kind="corporate_report", limit=3
        )
        self.assertTrue(report_chunks)
        self.assertTrue(all(chunk["source_locator"].get("page") for chunk in report_chunks))
        self.assertTrue(all(chunk["source_locator"]["filename"].lower().endswith(".pdf")
                            for chunk in report_chunks))
        report_source = provider.get_source_record(report_chunks[0]["source_record_id"])
        page_text = report_source["raw_payload"]["text"]
        cited_chunk = report_chunks[0]
        self.assertEqual(
            page_text[cited_chunk["start_offset"]:cited_chunk["end_offset"]],
            cited_chunk["original_text"],
        )
        with self.driver.session() as session:
            taxonomy_counts = session.run("""
                MATCH (topic:VFCTag)
                WITH count(topic) AS topic_count
                MATCH (issue:VFCTag3)
                RETURN topic_count, count(issue) AS issue_count
            """).single()
        self.assertGreater(taxonomy_counts["topic_count"], 0)
        self.assertGreater(taxonomy_counts["issue_count"], 0)

    def test_scoped_sales_and_native_survey_citations_resolve(self) -> None:
        from openpyxl import load_workbook

        provider = Neo4jEvidenceProvider(self.driver)
        sales = provider.get_sales_facts(EvidenceScope(model="G60", region="RoW"), year=2030)
        self.assertEqual(sales["facts"], [])
        self.assertEqual(len(sales["possible_whole_line_proxy_candidates"]), 1)
        fact = sales["possible_whole_line_proxy_candidates"][0]
        self.assertEqual(fact["vehicle_line"], "G60_G68")
        self.assertEqual(fact["region"], "RoW")
        self.assertEqual(sales["powertrain_split"], "unknown")
        locator = fact["source_locator"]
        self.assertEqual(locator["column_header"], "volume_2030")
        workbook = load_workbook(REAL_DATA_DIR / locator["filename"], read_only=True, data_only=True)
        try:
            self.assertEqual(workbook[locator["sheet"]].cell(locator["row"], locator["column"]).value, fact["volume"])
        finally:
            workbook.close()

        study = provider.get_survey_facts(None, EvidenceScope(model="G68", region="CN"))
        self.assertTrue(study["facts"])
        self.assertTrue(all(fact["derivative"] == "G68" and fact["region"] == "CN" for fact in study["facts"]))
        fact = study["facts"][0]
        locator = fact["source_locator"]
        self.assertTrue(locator["column"])
        workbook = load_workbook(REAL_DATA_DIR / locator["filename"], read_only=True, data_only=True)
        try:
            self.assertEqual(workbook[locator["sheet"]].cell(locator["row"], locator["column"]).value, fact["value"])
        finally:
            workbook.close()

        native_study = provider.get_survey_facts(None, EvidenceScope(model="G70", region="US"))
        distributions = [fact for fact in native_study["facts"] if fact["kind"] == "survey_distribution"]
        self.assertTrue(distributions)
        for distribution in distributions:
            self.assertEqual(len(distribution["distribution"]), 7)
            self.assertEqual(len(distribution["cell_locators"]), 7)
            self.assertIsNone(distribution["respondent_n"])
            for cell_locator in distribution["cell_locators"]:
                record = provider.get_source_record(cell_locator["source_record_id"])
                source_cell = record["raw_payload"]["columns"][cell_locator["column"] - 1]["value"]
                self.assertEqual(source_cell, distribution["distribution"][cell_locator["response_category"]])

        unmapped = provider.get_survey_facts("HVAC", EvidenceScope(model="G70", region="US"))
        self.assertEqual(unmapped["mapping_status"], "unmapped")
        self.assertEqual(unmapped["facts"], [])

    def test_summaries_have_exact_membership_and_are_idempotent(self) -> None:
        import runpy

        analysis = runpy.run_path(str(REPO_ROOT / "pipeline" / "03_vfc_analysis.py"))
        builder = runpy.run_path(str(REPO_ROOT / "pipeline" / "02_build_kg.py"))
        with self.driver.session() as session:
            first_analysis = analysis["analyze"](session)
            first_relationships = builder["build_evidence_relationships"](session)
            first_snapshot = self._graph_snapshot()
            second_analysis = analysis["analyze"](session)
            second_relationships = builder["build_evidence_relationships"](session)
            second_snapshot = self._graph_snapshot()
            invalid_summaries = session.run("""
                MATCH (stats:VFCAnalysis)
                OPTIONAL MATCH (stats)-[link:SUMMARIZES]->(fb:CustomerFeedback)
                WITH stats, count(DISTINCT fb.id) AS members, count(link) AS links
                WHERE stats.distinct_comments <> members OR members <> links
                RETURN stats.id AS id, stats.distinct_comments AS expected, members, links
                LIMIT 5
            """).data()
            incorrect_scope = session.run("""
                MATCH (stats:VFCAnalysis)-[:SUMMARIZES]->(fb:CustomerFeedback)
                WHERE coalesce(stats.model, '') <> coalesce(fb.model, '')
                   OR coalesce(stats.region, '') <> coalesce(fb.region, '')
                   OR coalesce(stats.powertrain, '') <> coalesce(fb.powertrain, '')
                RETURN count(*) AS count
            """).single()["count"]

        self.assertEqual(first_analysis, second_analysis)
        self.assertGreater(first_analysis["topics"], 0)
        self.assertGreater(first_analysis["issues"], 0)
        self.assertEqual(first_relationships, second_relationships)
        self.assertEqual(first_relationships["formal_conflicts_created"], 0)
        self.assertEqual(first_snapshot, second_snapshot)
        self.assertEqual(invalid_summaries, [])
        self.assertEqual(incorrect_scope, 0)


if __name__ == "__main__":
    unittest.main()
